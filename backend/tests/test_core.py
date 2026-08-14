"""核心单元测试：安全 / 切分 / 解析 / 融合。"""
from __future__ import annotations

import pytest

from app.core.security import create_access_token, decode_token, hash_password, verify_password
from app.rag.chunkers import chunk_document
from app.rag.parsers import parse_file
from app.rag.retriever import _rrf_fuse


# ---------------- 安全 ----------------
class TestSecurity:
    def test_password_roundtrip(self):
        h = hash_password("s3cret!")
        assert h != "s3cret!"
        assert verify_password("s3cret!", h)
        assert not verify_password("wrong", h)

    def test_jwt_roundtrip(self):
        t = create_access_token(1, "alice", "admin", 2)
        p = decode_token(t)
        assert p["sub"] == "1" and p["role"] == "admin" and p["tenant_id"] == 2

    def test_jwt_invalid(self):
        assert decode_token("not.a.token") is None


# ---------------- 切分 ----------------
class TestChunker:
    def test_fixed_chunk_with_overlap(self):
        text = "行" * 1000
        chunks = chunk_document(text, chunk_size=200, chunk_overlap=40)
        assert len(chunks) >= 5
        assert all(len(c["text"]) <= 200 for c in chunks)
        # 重叠：相邻块有共同内容
        assert chunks[0]["text"][-40:] in chunks[1]["text"]

    def test_markdown_heading_structure(self):
        md = "# 第一章\n\n企业简介内容……\n\n## 1.1 业务范围\n\n我们提供 SaaS 服务。\n\n# 第二章\n\n合规要求内容。\n"
        # 小块大小强制触发多块切分，验证标题保留
        chunks = chunk_document(md, chunk_size=30, chunk_overlap=0, is_markdown=True)
        assert len(chunks) >= 2
        assert any(c["metadata"].get("section") == "第一章" for c in chunks)

    def test_empty(self):
        assert chunk_document("   \n\n ") == []


# ---------------- 解析 ----------------
class TestParsers:
    def test_txt_utf8(self, tmp_path):
        p = tmp_path / "a.txt"
        p.write_text("你好，世界\n第二行", encoding="utf-8")
        parts = parse_file(p)
        assert parts[0][0] == "你好，世界\n第二行"
        assert parts[0][1] is None

    def test_gbk_fallback(self, tmp_path):
        p = tmp_path / "b.txt"
        p.write_bytes("中文内容".encode("gb18030"))
        parts = parse_file(p)
        assert "中文内容" in parts[0][0]

    def test_md(self, tmp_path):
        p = tmp_path / "c.md"
        p.write_text("# 标题\n正文", encoding="utf-8")
        assert parse_file(p)[0][0].startswith("# 标题")

    def test_docx(self, tmp_path):
        from docx import Document as Docx

        p = tmp_path / "d.docx"
        d = Docx()
        d.add_paragraph("这是 docx 段落一")
        d.add_paragraph("段落二")
        d.save(str(p))
        text = parse_file(p)[0][0]
        assert "docx 段落一" in text and "段落二" in text

    def test_xlsx(self, tmp_path):
        from openpyxl import Workbook

        p = tmp_path / "e.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.title = "数据表"
        ws.append(["姓名", "部门"])
        ws.append(["张三", "技术部"])
        wb.save(str(p))
        text = parse_file(p)[0][0]
        assert "张三" in text and "技术部" in text

    def test_unsupported(self, tmp_path):
        p = tmp_path / "f.exe"
        p.write_bytes(b"MZ")
        from app.rag.parsers import ParseError

        with pytest.raises(ParseError):
            parse_file(p)


# ---------------- RRF 融合 ----------------
class TestRRF:
    def test_fusion(self):
        vec = [{"chunk_id": 1, "score": 0.9}, {"chunk_id": 2, "score": 0.8}]
        kw = [{"chunk_id": 2, "score": 5.0}, {"chunk_id": 3, "score": 4.0}]
        out = _rrf_fuse([vec, kw], 3)
        ids = [o["chunk_id"] for o in out]
        assert ids == [2, 1, 3]  # 2 在两个来源中都出现 → 最高

    def test_empty(self):
        assert _rrf_fuse([[], []], 5) == []
